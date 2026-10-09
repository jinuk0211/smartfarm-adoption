"""Audit recorded regional crop-cycle sales, without treating sales as profit."""

from __future__ import annotations

import csv
import hashlib
import json
import math
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
CROPS = {"딸기", "오이", "토마토", "방울토마토", "파프리카"}
HEADERS = ["시설ID", "작기명", "정식일자", "작기종료일자", "품목", "품종", "재배면적",
           "재식수량", "시설유형", "시도", "시군구", "총출하량", "특상등급출하량",
           "상등급출하량", "증등급출하량", "하등급출하량", "수입금액"]
FEATURES = ["crop", "cultivar", "province", "district", "facility_type",
            "planting_year", "planting_month"]
TARGET = "revenue_krw_per_cycle"
PYEONG_TO_M2 = 400 / 121
SOURCE_SHA256 = "edbd318d7313805a9d7c731ec19fe06150e8cf5ba56e227eeec3f10d24a8da7a"
DEFINITION_SHA256 = "c1a7d684d6eb5646659aa6cf7a754f272902fe7e80ad456a6fae1d6006af8547"


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def canonical(value) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def raw_value(value):
    return value.isoformat() if isinstance(value, (date, datetime)) else value


def as_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(value) if isinstance(value, str) else None
    except ValueError:
        return None


def numeric(value) -> tuple[float | None, str]:
    if value is None or value == "":
        return None, "missing"
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None, "invalid"
    if isinstance(value, bool) or not math.isfinite(number):
        return None, "invalid"
    return number, "positive" if number > 0 else "zero" if number == 0 else "negative"


def text(value) -> str:
    return str(value).strip() if value is not None else ""


def verified_xlsx(path: Path, acquisition: Path) -> list[dict]:
    metadata = json.loads(acquisition.read_text(encoding="utf-8"))
    if (digest(path) != SOURCE_SHA256 or metadata.get("sha256") != SOURCE_SHA256
            or path.stat().st_size != metadata.get("bytes") or not metadata.get("acquired")):
        raise ValueError("Regional XLSX acquisition integrity check failed")
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if workbook.sheetnames != ["Sheet1"]:
            raise ValueError("Unexpected regional workbook sheets")
        rows = workbook["Sheet1"].iter_rows(values_only=True)
        if list(next(rows)) != HEADERS:
            raise ValueError("Unexpected regional source columns")
        return [dict(zip(HEADERS, row, strict=True)) for row in rows]
    finally:
        workbook.close()


def api_metadata(directory: Path) -> tuple[dict, dict]:
    """Read only the integrity-checked API metadata, never its outcome tables."""
    audit = json.loads((directory / "smartfarm_api_audit.json").read_text(encoding="utf-8"))
    tables, hashes = {}, {}
    for name in ("catalog", "cycles"):
        filename = f"smartfarm_api_{name}.csv"
        path = directory / filename
        expected = audit["output_files"][filename]
        if digest(path) != expected["sha256"] or path.stat().st_size != expected["bytes"]:
            raise ValueError("API metadata integrity check failed")
        hashes[filename] = digest(path)
        with path.open(encoding="utf-8-sig", newline="") as stream:
            tables[name] = list(csv.DictReader(stream))
    catalogs = defaultdict(list)
    for row in tables["catalog"]:
        catalogs[(row["fcltyId"], row["crpsnSn"], row["fcltyYear"])].append(row)
    matches = defaultdict(list)
    for row in tables["cycles"]:
        cats = catalogs[(row["fcltyId"], row["crpsnSn"], row["fcltyYear"])]
        key = (row["fcltyId"], row["fixplntngDe"], row["crpsnEndDe"], row["itemCodeNm"])
        area, status = numeric(row["ctvtAr"])
        cat_area = numeric(cats[0]["ctvtAr"])[0] if len(cats) == 1 else None
        valid = (status == "positive" and cat_area is not None
                 and math.isclose(area, cat_area, rel_tol=1e-9, abs_tol=1e-6))
        matches[key].append({"area_m2": area if valid else None,
                             "cycle_id": row["source_record_id"],
                             "catalog_id": cats[0]["source_record_id"] if len(cats) == 1 else ""})
    return dict(matches), hashes


def normalize(source_rows: list[dict], api_matches: dict | None = None) -> list[dict]:
    api_matches = api_matches or {}
    rows, identical, cycles = [], defaultdict(list), defaultdict(list)
    for excel_row, raw in enumerate(source_rows, 2):
        raw_json = canonical({key: raw_value(value) for key, value in raw.items()})
        start, end = as_date(raw["정식일자"]), as_date(raw["작기종료일자"])
        revenue, revenue_status = numeric(raw["수입금액"])
        area, area_status = numeric(raw["재배면적"])
        facility, crop = text(raw["시설ID"]), text(raw["품목"])
        cycle_key = (facility, start.isoformat() if start else "", crop)
        row = {"source_excel_row": excel_row, "source_raw_record_json": raw_json,
               "source_cell_types_json": canonical({key: type(value).__name__ for key, value in raw.items()}),
               "source_row_id": f"regional:{SOURCE_SHA256[:12]}:{excel_row}",
               "cycle_id": hashlib.sha256(canonical(cycle_key).encode()).hexdigest()[:24],
               "facility_id": facility, "crop": crop, "cultivar": text(raw["품종"]),
               "province": text(raw["시도"]), "district": text(raw["시군구"]),
               "facility_type": text(raw["시설유형"]), "cycle_label": text(raw["작기명"]),
               "planting_date": start.isoformat() if start else "",
               "cycle_end_date_audit_only": end.isoformat() if end else "",
               "planting_year": start.year if start else None,
               "planting_month": start.month if start else None,
               "area_source_definition_pyeong": area,
               "area_m2_if_definition_correct_audit_only": area * PYEONG_TO_M2 if area is not None else None,
               "area_status": area_status, TARGET: revenue, "revenue_status": revenue_status,
               "shipment_count_source_definition": numeric(raw["총출하량"])[0],
               "cycle_duration_days_audit_only": (end - start).days if end and start else None,
               "large_revenue_review_flag": revenue is not None and revenue > 10_000_000_000,
               "api_area_m2_verified_metadata": None, "api_area_comparison": "unmatched",
               "api_cycle_source_id": "", "api_catalog_source_id": ""}
        matches = api_matches.get((facility, row["planting_date"], row["cycle_end_date_audit_only"], crop), [])
        if len(matches) > 1:
            row["api_area_comparison"] = "ambiguous_cycle_match"
        elif len(matches) == 1:
            match = matches[0]
            row["api_cycle_source_id"], row["api_catalog_source_id"] = match["cycle_id"], match["catalog_id"]
            api_area = match["area_m2"]
            row["api_area_m2_verified_metadata"] = api_area
            row["api_area_comparison"] = "invalid_api_area" if api_area is None else "different_or_missing"
            if api_area is not None and area_status == "positive":
                if math.isclose(area, api_area, rel_tol=1e-6, abs_tol=0.01):
                    row["api_area_comparison"] = "raw_number_equals_api_m2_definition_conflict"
                elif math.isclose(area * PYEONG_TO_M2, api_area, rel_tol=1e-6, abs_tol=0.01):
                    row["api_area_comparison"] = "definition_conversion_agrees_api_m2"
        reasons = []
        if crop not in CROPS:
            reasons.append("crop_outside_five_crop_scope")
        if not facility:
            reasons.append("missing_facility")
        if not start or start.year < 1900:
            reasons.append("invalid_planting_date")
        if not end or end.year < 1900 or (start and end <= start):
            reasons.append("missing_or_invalid_completed_cycle_dates")
        if revenue_status != "positive":
            reasons.append(f"revenue_{revenue_status}")
        row["exclusion_reasons"] = reasons
        rows.append(row)
        identical[raw_json].append(row)
        cycles[cycle_key].append(row)
    for members in identical.values():
        source_rows_json = canonical([row["source_excel_row"] for row in members])
        for position, row in enumerate(members):
            row["exact_duplicate_source_rows_json"] = source_rows_json
            row["is_exact_duplicate_representative"] = position == 0
            if position:
                row["exclusion_reasons"].append("exact_duplicate_nonrepresentative")
    for members in cycles.values():
        # No supplied stable cycle ID: multiple distinct records on one facility,
        # planting date and crop are ambiguous, even if their labels/end dates differ.
        conflicting = len({row["source_raw_record_json"] for row in members}) > 1
        for row in members:
            row["ambiguous_cycle"] = conflicting
            if conflicting:
                row["exclusion_reasons"].append("ambiguous_facility_planting_crop_cycle")
            row["eligible_recorded_sales"] = not row["exclusion_reasons"]
            row["exclusion_reasons_json"] = canonical(row.pop("exclusion_reasons"))
    return rows


def write_csv(path: Path, rows: list[dict], fields: list[str] | None = None):
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields or list(rows[0]), extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def main():
    source = ROOT / "data/raw/smartfarm_curation/Regional_Crop_Producer_Shipping_Info.xlsx"
    acquisition = source.with_name("regional_shipping_acquisition.json")
    definition = ROOT / "data/raw/smartfarm_curation_docs/Regional_Crop_Shipping_Definition.hwp"
    if digest(definition) != DEFINITION_SHA256:
        raise ValueError("Regional definition integrity check failed")
    processed = ROOT / "data/processed"
    matches, api_hashes = api_metadata(processed)
    raw = verified_xlsx(source, acquisition)
    rows = normalize(raw, matches)
    candidates = [row for row in rows if row["eligible_recorded_sales"]]
    fields = ["source_row_id", "source_excel_row", "cycle_id", "facility_id", *FEATURES,
              TARGET, "api_area_m2_verified_metadata", "api_area_comparison",
              "large_revenue_review_flag", "exact_duplicate_source_rows_json"]
    paths = {"regional_curation_records.csv": rows,
             "regional_curation_candidates.csv": candidates}
    for name, contents in paths.items():
        write_csv(processed / name, contents, fields if "candidates" in name else None)
    statistics = {
        "source_rows": len(rows), "source_facilities": len({row["facility_id"] for row in rows}),
        "source_crop_counts": dict(Counter(row["crop"] for row in rows)),
        "revenue_status_counts": dict(Counter(row["revenue_status"] for row in rows)),
        "area_status_counts": dict(Counter(row["area_status"] for row in rows)),
        "raw_cell_type_counts": {name: dict(Counter(type(row[name]).__name__ for row in raw)) for name in HEADERS},
        "missing_cell_counts": {name: sum(row[name] is None for row in raw) for name in HEADERS},
        "exact_duplicate_nonrepresentative_rows": sum(not row["is_exact_duplicate_representative"] for row in rows),
        "ambiguous_cycle_rows": sum(row["ambiguous_cycle"] for row in rows),
        "ambiguous_positive_sales_rows": sum(row["ambiguous_cycle"] and row["revenue_status"] == "positive" for row in rows),
        "exclusion_reason_counts": dict(Counter(reason for row in rows for reason in json.loads(row["exclusion_reasons_json"]))),
        "eligible_rows": len(candidates), "eligible_facilities": len({row["facility_id"] for row in candidates}),
        "eligible_crop_counts": dict(Counter(row["crop"] for row in candidates)),
        "eligible_planting_year_counts": dict(Counter(row["planting_year"] for row in candidates)),
        "api_area_comparison_all_rows": dict(Counter(row["api_area_comparison"] for row in rows)),
        "api_area_comparison_eligible_rows": dict(Counter(row["api_area_comparison"] for row in candidates)),
        "eligible_verified_api_area_rows": sum(row["api_area_m2_verified_metadata"] is not None for row in candidates),
        "large_revenue_review_rows": [{key: row[key] for key in ("source_row_id", "facility_id", "crop", TARGET)} for row in candidates if row["large_revenue_review_flag"]],
        "eligible_min_revenue_krw": min(row[TARGET] for row in candidates),
        "eligible_max_revenue_krw": max(row[TARGET] for row in candidates),
    }
    manifest = {"prepared_at_utc": datetime.now(timezone.utc).isoformat(),
                "source": str(source.relative_to(ROOT)), "source_sha256": digest(source),
                "definition_sha256": digest(definition), "acquisition_sha256": digest(acquisition),
                "preparation_code_sha256": digest(Path(__file__)), "api_metadata_sha256": api_hashes,
                "target": TARGET, "target_definition": "Recorded crop-cycle gross sales in KRW; not profit or causal adoption effect",
                "primary_feature_whitelist": FEATURES, "group_column": "facility_id",
                "source_units": {"area": "pyeong (definition; cross-source contradiction observed)", "shipment": "count", "sales": "KRW"},
                "area_policy": "No regional-source area input in primary model. Conditional definition conversion is audit-only; exact unique API m2 subset may be evaluated separately.",
                "outlier_policy": "All eligible recorded sales retained. Review flag > KRW 10 billion does not exclude/correct. Report explicit sensitivity separately.",
                "selection_policy": "Five crops, positive recorded sales, valid completed cycle dates, unambiguous facility+planting+crop identity; exact duplicates mapped to one representative.",
                "statistics": statistics,
                "outputs": {name: {"sha256": digest(processed / name), "bytes": (processed / name).stat().st_size} for name in paths}}
    (processed / "regional_curation_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    report = ["# 지역별 품목별 농가출하정보 — 실제 원자료 감사", "",
              f"검증된 XLSX {len(rows):,}행, {statistics['source_facilities']:,}개 시설. 원자료 SHA256 `{digest(source)}`.", "",
              "정의서는 총출하량을 작기별 **개**, 수입금액을 **총매출 원**, 면적을 **평**으로 정의한다. kg·순이익·스마트팜 도입 효과로 해석하지 않는다.", "",
              "## 모델용 계약", "",
              f"`regional_curation_candidates.csv`: {len(candidates)}개 작기, {statistics['eligible_facilities']}개 시설. 타깃 `{TARGET}`. 시설 ID로 그룹 분리한다.",
              f"주 입력 열: `{', '.join(FEATURES)}`. 정식 연·월은 사용자가 계획하는 시점으로 해석하며, 실제 종료일·작기 길이·출하량·등급·매출·감사 플래그는 입력에 넣지 않는다.",
              "품종·시설유형 공란은 원본 결측으로 보존한다. 모델 전처리는 학습 fold에서만 수행한다.", "",
              "면적은 정의서의 평과 정확히 일치하는 API 메타데이터의 ㎡가 충돌한다. 원 숫자가 API ㎡와 같은 레코드가 존재하므로 평→㎡ 변환값을 확정 면적으로 사용하지 않는다. 전 행 파일의 조건부 환산 열은 감사 전용이다. 기본 모델은 면적을 제외하고, 별도 실험에서만 시설+정식일+종료일+품목의 유일한 API 매칭과 API catalog/cycle 면적 일치를 통과한 ㎡를 사용할 수 있다.", "",
              "## 선별과 한계", "",
              "목표 5개 품목, 양의 기록 매출, 유효 종료일, 시설+정식일+품목 키가 모호하지 않은 행만 후보로 사용한다. 작기 이름과 종료일이 다른 중복도 동일 식별키에 여러 원문이 있으면 모두 제외한다. 정확히 동일한 행은 한 대표행과 모든 Excel 행번호를 보존한다. 합산·임의 보정은 하지 않는다.",
              "0과 결측은 구별하며 결측 매출을 0으로 채우지 않는다. 기록이 없는 농가가 다수이므로 성과가 기록된 표본에 대한 선택 편향이 있다. 실제 종료일 검사는 작기 전체 기록의 독립적 검증이 아니다. 매출 정의에 근거한 기록값 모델이며 미래 농가 수입을 보증하지 않는다.",
              "100억 원 초과 검토 플래그는 원자료를 삭제하거나 교정하지 않는다. 원값 포함 기본 결과와 해당 플래그를 명시한 민감도 결과를 구분해야 한다. 비용이 없어 순이익·ROI 정답을 만들 수 없다.", "",
              "## 실제 감사 수치", "", "```json", json.dumps(statistics, ensure_ascii=False, indent=2), "```", "",
              "원문 17개 필드(원본 오타 `증등급출하량` 포함), 셀 타입, Excel 행, 원시 행 JSON, 중복 매핑, 제외 사유는 `regional_curation_records.csv`에 보존한다. 출처·정의서·API 메타데이터·스크립트·결과 해시는 manifest에 기록한다."]
    (ROOT / "reports/regional_curation_schema_audit.md").write_text("\n".join(report) + "\n", encoding="utf-8")
    print(json.dumps(statistics, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
